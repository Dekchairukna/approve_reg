"""นำเข้ารายชื่อนักกีฬาจาก Excel

ต้องมีแถวหัวตารางที่มีคอลัมน์อย่างน้อย: รายการแข่งขัน, โรงเรียน, ชื่อ, นามสกุล
คอลัมน์เสริม: คำนำหน้า, ระดับชั้น, คู่ที่/ทีมที่/มือวาง
"""
import re

from openpyxl import load_workbook

LEVEL_ORDER = {'อนุบาล': 0, 'ประถมศึกษา': 1, 'มัธยมศึกษาตอนต้น': 2, 'มัธยมศึกษาตอนปลาย': 3}
CATEGORY_RULES = [('คู่ผสม', 'คู่ผสม'), ('Shooting', 'Shooting'), ('เดี่ยว', 'เดี่ยว'), ('บุคคล', 'เดี่ยว'),
                  ('ทีม', 'ทีม'), ('คู่', 'คู่')]
CATEGORY_ORDER = {'เดี่ยว': 0, 'Shooting': 1, 'คู่ผสม': 2, 'คู่': 3, 'ทีม': 4}


class ImportError_(Exception):
    pass


def _find_columns(header):
    cols = {}
    for i, h in enumerate(header):
        h = re.sub(r'\s+', '', str(h or ''))
        if not h:
            continue
        if 'event' not in cols and ('รายการ' in h or h.startswith('ประเภท')):
            cols['event'] = i
        elif 'school' not in cols and ('โรงเรียน' in h or 'สถาบัน' in h or 'สังกัด' in h):
            cols['school'] = i
        elif 'prefix' not in cols and 'คำนำหน้า' in h:
            cols['prefix'] = i
        elif 'last' not in cols and 'นามสกุล' in h:
            cols['last'] = i
        elif 'first' not in cols and h.startswith('ชื่อ') and 'ทีม' not in h:
            cols['first'] = i
        elif 'level' not in cols and 'ระดับ' in h:
            cols['level'] = i
        elif 'team' not in cols and ('คู่ที่' in h or 'ทีมที่' in h or 'มือวาง' in h):
            cols['team'] = i
    return cols


def read_excel(path_or_file):
    try:
        wb = load_workbook(path_or_file, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise ImportError_(f'อ่านไฟล์ Excel ไม่ได้: {e}')
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    for hi, header in enumerate(rows[:20]):
        cols = _find_columns(header)
        if {'event', 'school', 'first', 'last'} <= cols.keys():
            break
    else:
        raise ImportError_('ไม่พบหัวตาราง ต้องมีคอลัมน์ รายการแข่งขัน, โรงเรียน, ชื่อ, นามสกุล')

    def cell(r, k):
        i = cols.get(k)
        v = r[i] if i is not None and i < len(r) else ''
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        return re.sub(r'\s+', ' ', str(v if v is not None else '')).strip()

    out = []
    for r in rows[hi + 1:]:
        if not r or not cell(r, 'event') or not cell(r, 'first'):
            continue
        out.append(dict(event=cell(r, 'event'), school=cell(r, 'school'), prefix=cell(r, 'prefix'),
                        first=cell(r, 'first'), last=cell(r, 'last'), level=cell(r, 'level'),
                        team=cell(r, 'team') or '999'))
    if not out:
        raise ImportError_('ไม่พบรายชื่อนักกีฬาในไฟล์')
    return out


def category_of(name):
    for key, cat in CATEGORY_RULES:
        if key in name:
            return cat
    return 'อื่นๆ'


def short_of(name, level):
    s = re.sub(r'^[^:]*:\s*', '', name)          # ตัด "เปตอง : "
    for lv in list(LEVEL_ORDER) + ([level] if level else []):
        s = s.replace(lv, '')
    s = re.sub(r'\s+', ' ', s.replace('บุคคล', '').replace('1 คน', '')).strip()
    s = re.sub(r'^(เดี่ยว|คู่) (ชาย|หญิง)$', r'\1\2', s)
    return s or name


def import_rows(conn, comp_id, rows):
    """แทนที่รายชื่อ/รายการ/การยืนยันทั้งหมดของการแข่งขันนี้ด้วยข้อมูลใหม่"""
    events = {}
    for r in rows:
        if r['event'] not in events:
            events[r['event']] = dict(name=r['event'], level=r['level'], category=category_of(r['event']),
                                      short=short_of(r['event'], r['level']))
    ordered = sorted(events.values(), key=lambda e: (CATEGORY_ORDER.get(e['category'], 9),
                                                      LEVEL_ORDER.get(e['level'], 9), e['short']))
    old = {r['name']: r for r in conn.execute(
        'SELECT name, start_date, start_time, deadline FROM events WHERE comp_id=?', (comp_id,))}
    conn.execute('DELETE FROM confs WHERE comp_id=?', (comp_id,))
    conn.execute('DELETE FROM entries WHERE comp_id=?', (comp_id,))
    conn.execute('DELETE FROM events WHERE comp_id=?', (comp_id,))
    for i, e in enumerate(ordered):
        o = old.get(e['name'])   # คงวันแข่ง/เวลาปิดเดิมไว้ถ้าชื่อรายการเหมือนเดิม
        conn.execute('INSERT INTO events (comp_id, name, short, level, category, start_date, start_time, deadline, sort) '
                     'VALUES (?,?,?,?,?,?,?,?,?)',
                     (comp_id, e['name'], e['short'], e['level'], e['category'],
                      o['start_date'] if o else '', o['start_time'] if o else '', o['deadline'] if o else '', i))
    conn.executemany(
        'INSERT INTO entries (comp_id, event, school, level, team, prefix, first, last, o_prefix, o_first, o_last) '
        'VALUES (?,?,?,?,?,?,?,?,?,?,?)',
        [(comp_id, r['event'], r['school'], r['level'], r['team'], r['prefix'], r['first'], r['last'],
          r['prefix'], r['first'], r['last']) for r in rows])
    return dict(entries=len(rows), events=len(ordered), schools=len({r['school'] for r in rows}))
