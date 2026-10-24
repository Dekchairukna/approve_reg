"""หลังบ้านเจ้าหน้าที่: จัดการการแข่งขัน วันแข่ง ภาพรวม พิมพ์ ส่งออก"""
import io
import os
import re
import uuid

from flask import Blueprint, abort, current_app, g, jsonify, request, send_file

from . import logic, store
from .auth import err, login_required
from .db import get_db
from .importer import ImportError_, import_rows, read_excel

bp = Blueprint('admin', __name__)
LEVEL_SHORT = {'ประถมศึกษา': 'ประถม', 'มัธยมศึกษาตอนต้น': 'ม.ต้น', 'มัธยมศึกษาตอนปลาย': 'ม.ปลาย'}
LOGO_TYPES = {'.png': 'png', '.jpg': 'jpg', '.jpeg': 'jpg', '.webp': 'webp', '.gif': 'gif'}


@bp.get('/admin')
def admin_page():
    from flask import send_from_directory
    return send_from_directory(current_app.static_folder, 'admin.html')


def _comp_or_404(cid):
    c = store.comp_by_id(cid)
    if not c:
        abort(404)
    return c


def _comp_summary(c):
    conn = get_db()
    d = store.comp_public(c)
    d.update(
        entries=conn.execute('SELECT COUNT(*) FROM entries WHERE comp_id=?', (c['id'],)).fetchone()[0],
        schools=conn.execute('SELECT COUNT(DISTINCT school) FROM entries WHERE comp_id=?', (c['id'],)).fetchone()[0],
        events=conn.execute('SELECT COUNT(*) FROM events WHERE comp_id=?', (c['id'],)).fetchone()[0],
        confirmed=conn.execute('SELECT COUNT(*) FROM confs WHERE comp_id=?', (c['id'],)).fetchone()[0],
        undated=conn.execute("SELECT COUNT(*) FROM events WHERE comp_id=? AND deadline=''", (c['id'],)).fetchone()[0],
        created_at=c['created_at'])
    return d


def _read_logo(file):
    """ตรวจไฟล์โลโก้ คืน (ชนิดไฟล์, ข้อมูล) — เก็บในฐานข้อมูล เพื่อไม่หายเมื่อ deploy ใหม่"""
    ext = os.path.splitext(file.filename or '')[1].lower()
    if ext not in LOGO_TYPES:
        raise ValueError('โลโก้ต้องเป็นไฟล์ png, jpg, webp หรือ gif')
    data = file.read()
    if len(data) > 3 * 1024 * 1024:
        raise ValueError('ไฟล์โลโก้ใหญ่เกิน 3 MB')
    return 'image/' + ('jpeg' if LOGO_TYPES[ext] == 'jpg' else LOGO_TYPES[ext]), data


def _valid_slug(s):
    return re.fullmatch(r'[a-z0-9][a-z0-9-]{1,40}', s or '') is not None


# ---------- การแข่งขัน ----------
@bp.get('/api/admin/competitions')
@login_required()
def list_comps():
    rows = get_db().execute(f'SELECT {store.COMP_COLS} FROM competitions ORDER BY active DESC, id DESC')
    return jsonify(competitions=[_comp_summary(dict(r)) for r in rows], me=g.user)


@bp.post('/api/admin/competitions')
@login_required('admin')
def create_comp():
    f = request.form
    name, slug = (f.get('name') or '').strip(), (f.get('slug') or '').strip().lower()
    if not name:
        return err('กรุณากรอกชื่อการแข่งขัน')
    if not _valid_slug(slug):
        return err('รหัสลิงก์ต้องเป็นอังกฤษตัวเล็ก/ตัวเลข/ขีด 2–41 ตัว เช่น petong49')
    if 'file' not in request.files or not request.files['file'].filename:
        return err('กรุณาเลือกไฟล์ Excel รายชื่อนักกีฬา')
    try:
        max_events = max(0, int(f.get('max_events') or 3))
    except ValueError:
        return err('จำนวนรายการสูงสุดต่อคนต้องเป็นตัวเลข')
    conn = get_db()
    if conn.execute('SELECT 1 FROM competitions WHERE slug=?', (slug,)).fetchone():
        return err('รหัสลิงก์นี้ถูกใช้แล้ว')
    try:
        rows = read_excel(request.files['file'])
        logo = _read_logo(request.files['logo']) if request.files.get('logo') and request.files['logo'].filename else None
    except (ImportError_, ValueError) as e:
        return err(str(e))
    with conn:
        cid = conn.insert_id('INSERT INTO competitions (slug, name, sport, subtitle, max_events, created_at) '
                             'VALUES (?,?,?,?,?,?)',
                             (slug, name, (f.get('sport') or '').strip(), (f.get('subtitle') or '').strip(),
                              max_events, logic.stamp()))
        result = import_rows(conn, cid, rows)
        if logo:
            conn.execute('UPDATE competitions SET logo=?, logo_mime=?, logo_data=? WHERE id=?',
                         (uuid.uuid4().hex[:8], logo[0], logo[1], cid))
    return jsonify(ok=True, id=cid, imported=result)


@bp.get('/api/admin/competitions/<int:cid>')
@login_required()
def get_comp(cid):
    c = _comp_or_404(cid)
    return jsonify(comp=_comp_summary(c), events=store.load_events(cid), schools=store.schools(cid), me=g.user)


@bp.post('/api/admin/competitions/<int:cid>')
@login_required('admin')
def update_comp(cid):
    c = _comp_or_404(cid)
    f = request.form if request.form else (request.get_json(silent=True) or {})
    slug = (f.get('slug') or c['slug']).strip().lower()
    if not _valid_slug(slug):
        return err('รหัสลิงก์ไม่ถูกต้อง')
    conn = get_db()
    if conn.execute('SELECT 1 FROM competitions WHERE slug=? AND id<>?', (slug, cid)).fetchone():
        return err('รหัสลิงก์นี้ถูกใช้แล้ว')
    try:
        max_events = max(0, int(f.get('max_events', c['max_events'])))
    except ValueError:
        return err('จำนวนรายการสูงสุดต่อคนต้องเป็นตัวเลข')
    active = str(f.get('active', c['active'])).lower() in ('1', 'true', 'on')
    logo = None
    if request.files.get('logo') and request.files['logo'].filename:
        try:
            logo = _read_logo(request.files['logo'])
        except ValueError as e:
            return err(str(e))
    with conn:
        conn.execute('UPDATE competitions SET slug=?, name=?, sport=?, subtitle=?, max_events=?, active=? WHERE id=?',
                     (slug, (f.get('name') or c['name']).strip(), (f.get('sport', c['sport']) or '').strip(),
                      (f.get('subtitle', c['subtitle']) or '').strip(), max_events, int(active), cid))
        if logo:
            conn.execute('UPDATE competitions SET logo=?, logo_mime=?, logo_data=? WHERE id=?',
                         (uuid.uuid4().hex[:8], logo[0], logo[1], cid))
    return jsonify(ok=True)


@bp.post('/api/admin/competitions/<int:cid>/delete')
@login_required('admin')
def delete_comp(cid):
    c = _comp_or_404(cid)
    if ((request.get_json(silent=True) or {}).get('confirm') or '').strip().lower() != c['slug']:
        return err(f"พิมพ์รหัสลิงก์ \"{c['slug']}\" ให้ตรงเพื่อยืนยันการลบ")
    conn = get_db()
    with conn:
        for t in ('logs', 'confs', 'entries', 'events'):
            conn.execute(f'DELETE FROM {t} WHERE comp_id=?', (cid,))
        conn.execute('DELETE FROM competitions WHERE id=?', (cid,))
    return jsonify(ok=True)


@bp.post('/api/admin/competitions/<int:cid>/import')
@login_required('admin')
def reimport(cid):
    _comp_or_404(cid)
    conn = get_db()
    confirmed = conn.execute('SELECT COUNT(*) FROM confs WHERE comp_id=?', (cid,)).fetchone()[0]
    if confirmed and request.form.get('confirm') != 'yes':
        return err(f'มีการยืนยันไปแล้ว {confirmed} รายการ การนำเข้าใหม่จะลบการยืนยันทั้งหมด', 409)
    if 'file' not in request.files or not request.files['file'].filename:
        return err('กรุณาเลือกไฟล์ Excel')
    try:
        rows = read_excel(request.files['file'])
    except ImportError_ as e:
        return err(str(e))
    with conn:
        result = import_rows(conn, cid, rows)
        conn.execute('INSERT INTO logs (comp_id, at, action, after, staff_user) VALUES (?,?,?,?,?)',
                     (cid, logic.stamp(), 'นำเข้ารายชื่อใหม่', f"{result['entries']} รายชื่อ", g.user['username']))
    return jsonify(ok=True, imported=result)


@bp.post('/api/admin/competitions/<int:cid>/events')
@login_required('admin')
def update_events(cid):
    _comp_or_404(cid)
    items = (request.get_json(silent=True) or {}).get('events') or []
    errors = []
    for it in items:
        if it.get('deadline') and not logic.parse_deadline(it['deadline']):
            errors.append(f"เวลาปิดยืนยันไม่ถูกต้อง: {it.get('short') or it.get('id')}")
        if it.get('startDate') and not re.fullmatch(r'\d{4}-\d{2}-\d{2}', it['startDate']):
            errors.append(f"วันแข่งไม่ถูกต้อง: {it.get('short') or it.get('id')}")
    if errors:
        return err(' · '.join(errors))
    conn = get_db()
    with conn:
        for it in items:
            conn.execute('UPDATE events SET short=?, start_date=?, start_time=?, deadline=? WHERE id=? AND comp_id=?',
                         ((it.get('short') or '').strip(), it.get('startDate') or '', (it.get('startTime') or '').strip(),
                          (it.get('deadline') or '').replace('T', ' ')[:16], int(it['id']), cid))
    return jsonify(ok=True, events=store.load_events(cid))


@bp.get('/api/admin/competitions/<int:cid>/overview')
@login_required()
def overview(cid):
    c = _comp_or_404(cid)
    entries, events, confs = store.load_all(cid)
    return jsonify(logic.overview(entries, events, confs, logic.now_bkk(), c['max_events']))


@bp.post('/api/admin/competitions/<int:cid>/print')
@login_required()
def print_data(cid):
    _comp_or_404(cid)
    entries, events, confs = store.load_all(cid)
    return jsonify(logic.print_data(entries, events, confs, request.get_json(silent=True) or {}))


@bp.get('/api/admin/competitions/<int:cid>/checkin')
@login_required()
def checkin_list(cid):
    """รายชื่อสำหรับรับรายงานตัว: นักกีฬาที่ยืนยันแล้วของรายการที่เลือก (เลือกได้หลายรายการ ?event=..&event=..)"""
    _comp_or_404(cid)
    names = [n for n in request.args.getlist('event') if n]
    em = {e['event']: e for e in store.load_events(cid)}
    names = [n for n in names if n in em]
    if not names:
        return jsonify(rows=[], unconfirmed=[], events=[])
    want = set(names)
    all_entries, _, all_confs = store.load_all(cid)
    entries = [e for e in all_entries if e['event'] in want]
    confirmed = {(c['school'], c['event']) for c in all_confs if c['event'] in want}
    rows = sorted((e for e in entries if e['status'] == logic.OK),
                  key=lambda e: (e['school'], e['first'], e['last'], em[e['event']].get('sort', 0)))
    out = [dict(id=e['id'], school=e['school'], event=e['event'], short=em[e['event']]['short'],
                level=em[e['event']]['level'], teamLabel=logic.team_label(e['team'], em[e['event']]['category']),
                name=f"{e['prefix']}{e['first']} {e['last']}", key=logic.pkey(e['school'], e['first'], e['last']),
                checkinAt=e['checkinAt'], checkinBy=e['checkinBy'], auto=bool(e.get('auto'))) for e in rows]
    unconfirmed = sorted({(e['school'], em[e['event']]['short'], em[e['event']]['level']) for e in entries
                          if (e['school'], e['event']) not in confirmed})
    return jsonify(rows=out, events=names,
                   unconfirmed=[dict(school=s, short=sh, level=lv) for s, sh, lv in unconfirmed])


@bp.get('/api/admin/competitions/<int:cid>/checkin-report.xlsx')
@login_required()
def checkin_report(cid):
    """รายงานการรายงานตัว แยกตามรายการ: ชื่อ โรงเรียน ทีม เวลา (มาแล้วเรียงตามเวลา ตามด้วยยังไม่มา)"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    c = _comp_or_404(cid)
    want = set(request.args.getlist('event'))
    entries, events, _ = store.load_all(cid)
    wb = Workbook()
    ws = wb.active
    ws.title = 'รายงานตัว'
    ws.append([c['name']])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([f'พิมพ์เมื่อ {logic.now_bkk():%Y-%m-%d %H:%M} น.'])
    head = ['ที่', 'ชื่อ-นามสกุล', 'โรงเรียน', 'ทีม/คู่', 'เวลารายงานตัว', 'วันที่', 'ผู้รับรายงานตัว']
    for ev in events:
        if want and ev['event'] not in want:
            continue
        rows = [e for e in entries if e['event'] == ev['event'] and e['status'] == logic.OK]
        if not rows:
            continue
        rows.sort(key=lambda e: (0, e['checkinAt']) if e['checkinAt'] else (1, e['school'] + e['first'] + e['last']))
        n = sum(1 for e in rows if e['checkinAt'])
        ws.append([])
        ws.append([f"{LEVEL_SHORT.get(ev['level'], ev['level'])} · {ev['short']}  ({ev['event']})"
                   + (f"  แข่ง {ev['startDate']}" if ev['startDate'] else '') + f"  —  มา {n}/{len(rows)} คน"])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=12)
        ws.append(head)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='EEEEEE')
        for i, e in enumerate(rows, 1):
            at = e['checkinAt'] or ''
            ws.append([i, f"{e['prefix']}{e['first']} {e['last']}", e['school'], logic.team_label(e['team'], ev['category']),
                       at[11:16] if at else 'ยังไม่มา', at[:10], e['checkinBy'] or ''])
    for col, w in zip('ABCDEFG', (6, 32, 40, 14, 14, 12, 18)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"รายงานตัว_{c['slug']}_{logic.now_bkk():%Y%m%d_%H%M}.xlsx",
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@bp.post('/api/admin/competitions/<int:cid>/checkin')
@login_required()
def checkin_mark(cid):
    """บันทึกรายงานตัว ส่ง ids (หลายรายการของเด็กคนเดียว) หรือ id"""
    _comp_or_404(cid)
    d = request.get_json(silent=True) or {}
    try:
        ids = [int(x) for x in (d.get('ids') or [d.get('id')])]
    except (TypeError, ValueError):
        return err('ข้อมูลไม่ถูกต้อง')
    conn = get_db()
    present = bool(d.get('present'))
    at = logic.stamp() if present else ''
    who = (g.user.get('name') or g.user['username']) if present else ''
    done = []
    with conn:
        for eid in ids:
            row = conn.execute('SELECT * FROM entries WHERE id=? AND comp_id=?', (eid, cid)).fetchone()
            if not row:
                continue
            conn.execute('UPDATE entries SET checkin_at=?, checkin_by=? WHERE id=?', (at, who, eid))
            conn.execute('INSERT INTO logs (comp_id, at, school, event, entry_id, action, before, after, staff_user) '
                         'VALUES (?,?,?,?,?,?,?,?,?)',
                         (cid, logic.stamp(), row['school'], row['event'], eid,
                          'รายงานตัว' if present else 'ยกเลิกรายงานตัว', '',
                          f"{row['prefix']}{row['first']} {row['last']}", g.user['username']))
            done.append(eid)
    if not done:
        return err('ไม่พบนักกีฬา', 404)
    return jsonify(ok=True, ids=done, checkinAt=at, checkinBy=who)


@bp.get('/api/admin/competitions/<int:cid>/roster.xlsx')
@login_required()
def roster_excel(cid):
    """รายชื่อที่ยืนยันแล้ว (รวมยืนยันอัตโนมัติ) แบบตาราง สำหรับนำไปลงระบบอื่น เช่น ระบบ Shooting
    กรองได้ด้วย ?startDate=&level=&category=&school="""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    c = _comp_or_404(cid)
    f = {k: request.args.get(k, '') for k in ('startDate', 'level', 'category', 'school')}
    entries, events, confs = store.load_all(cid)
    d = logic.print_data(entries, events, confs, f)
    em = {e['event']: e for e in events}
    by_key = {}
    for e in entries:
        if e['status'] == logic.OK:
            by_key.setdefault((e['school'], e['event']), []).append(e)
    wb = Workbook()
    ws = wb.active
    ws.title = 'รายชื่อที่ยืนยัน'
    head = ['ลำดับ', 'รายการแข่งขัน', 'ชื่อย่อรายการ', 'ประเภท', 'ระดับชั้น', 'วันแข่ง', 'โรงเรียน', 'ทีม/คู่',
            'คำนำหน้า', 'ชื่อ', 'นามสกุล', 'ชื่อ-นามสกุล', 'การยืนยัน']
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    n = 0
    for page in d['pages']:
        for b in page['blocks']:
            ev = em[b['event']]
            for e in sorted(by_key.get((page['school'], b['event']), []), key=logic.team_sort):
                n += 1
                ws.append([n, b['event'], ev['short'], ev['category'], ev['level'], ev['startDate'], page['school'],
                           logic.team_label(e['team'], ev['category']), e['prefix'], e['first'], e['last'],
                           f"{e['prefix']}{e['first']} {e['last']}",
                           'อัตโนมัติ (เลยเวลาปิด)' if b['conf'].get('auto') else 'โรงเรียนยืนยัน'])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(max(len(str(x.value or '')) for x in col) + 2, 50)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    tag = '_'.join(x for x in (f['category'], f['startDate']) if x) or 'ทั้งหมด'
    return send_file(buf, as_attachment=True, download_name=f"รายชื่อยืนยัน_{c['slug']}_{tag}.xlsx",
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@bp.get('/api/admin/competitions/<int:cid>/athletes')
@login_required()
def athletes(cid):
    """รายชื่อนักกีฬา (ไม่ซ้ำคน) ของโรงเรียน สำหรับแก้ชื่อที่สะกดผิด"""
    _comp_or_404(cid)
    school = request.args.get('school') or ''
    em = {e['event']: e for e in store.load_events(cid)}
    people = {}
    for e in store.load_entries(cid, school):
        k = logic.pkey(e['school'], e['first'], e['last'])
        p = people.setdefault(k, dict(key=k, prefix=e['prefix'], first=e['first'], last=e['last'], events=[]))
        if e['status'] != logic.OUT:
            ev = em.get(e['event'], {})
            lv = {'ประถมศึกษา': 'ประถม', 'มัธยมศึกษาตอนต้น': 'ม.ต้น', 'มัธยมศึกษาตอนปลาย': 'ม.ปลาย'}.get(ev.get('level', ''), ev.get('level', ''))
            p['events'].append(f"{lv} {ev.get('short', e['event'])}".strip())
    rows = sorted(people.values(), key=lambda p: (p['first'], p['last']))
    return jsonify(athletes=rows)


@bp.post('/api/admin/competitions/<int:cid>/athletes/rename')
@login_required('admin')
def rename_athlete(cid):
    """แก้ชื่อสะกดผิด: เปลี่ยนทุกแถวของเด็กคนนี้ในโรงเรียนนี้ ทั้งชื่อปัจจุบันและชื่อตามใบสมัคร"""
    _comp_or_404(cid)
    d = request.get_json(silent=True) or {}
    school, key = d.get('school') or '', d.get('key') or ''
    p, f, l = logic.clean(d.get('prefix'))[:20], logic.clean(d.get('first'))[:80], logic.clean(d.get('last'))[:80]
    if not f or not l:
        return err('กรุณากรอกชื่อและนามสกุล')
    conn = get_db()
    rows = conn.execute('SELECT * FROM entries WHERE comp_id=? AND school=?', (cid, school)).fetchall()
    cur = [r for r in rows if logic.pkey(r['school'], r['first'], r['last']) == key]
    orig = [r for r in rows if logic.pkey(r['school'], r['o_first'], r['o_last']) == key]
    if not cur and not orig:
        return err('ไม่พบนักกีฬาคนนี้ กรุณาโหลดหน้าใหม่', 404)
    old = cur[0] if cur else orig[0]
    old_name = f"{old['prefix']}{old['first']} {old['last']}"
    with conn:
        for r in cur:
            conn.execute('UPDATE entries SET prefix=?, first=?, last=? WHERE id=?', (p, f, l, r['id']))
        for r in orig:
            conn.execute('UPDATE entries SET o_prefix=?, o_first=?, o_last=? WHERE id=?', (p, f, l, r['id']))
        conn.execute('INSERT INTO logs (comp_id, at, school, event, action, before, after, staff_user) VALUES (?,?,?,?,?,?,?,?)',
                     (cid, logic.stamp(), school, '', 'แก้ชื่อสะกด', old_name, f'{p}{f} {l}', g.user['username']))
    return jsonify(ok=True, updated=len({r['id'] for r in cur + orig}))


@bp.get('/api/admin/competitions/<int:cid>/logs')
@login_required()
def logs(cid):
    _comp_or_404(cid)
    rows = get_db().execute('SELECT * FROM logs WHERE comp_id=? ORDER BY id DESC LIMIT 300', (cid,))
    return jsonify(logs=[dict(r, by=r['by_name'], user=r['staff_user']) for r in rows])


@bp.get('/api/admin/backup')
@login_required('admin')
def backup():
    from .backup import backup_to_sqlite_bytes
    data = backup_to_sqlite_bytes(current_app.config['DATABASE'])
    return send_file(io.BytesIO(data), as_attachment=True, download_name=f'checkin_backup_{logic.now_bkk():%Y%m%d_%H%M}.db',
                     mimetype='application/octet-stream')


@bp.post('/api/admin/restore')
@login_required('admin')
def restore():
    from .backup import restore_from_sqlite_bytes
    if request.form.get('confirm') != 'yes':
        return err('ติ๊กยืนยันก่อน การกู้คืนจะแทนที่ข้อมูลทั้งหมดในระบบ')
    f = request.files.get('file')
    if not f:
        return err('กรุณาเลือกไฟล์สำรอง (.db)')
    try:
        counts = restore_from_sqlite_bytes(f.read(), current_app.config['DATABASE'])
    except ValueError as e:
        return err(str(e))
    return jsonify(ok=True, counts=counts)


@bp.get('/api/admin/competitions/<int:cid>/export')
@login_required()
def export_excel(cid):
    from openpyxl import Workbook
    c = _comp_or_404(cid)
    entries, events, confs = store.load_all(cid)
    em = {e['event']: e for e in events}
    wb = Workbook()
    ws = wb.active
    ws.title = 'รายชื่อนักกีฬา'
    ws.append(['รายการแข่งขัน', 'โรงเรียน', 'คำนำหน้า', 'ชื่อ', 'นามสกุล', 'ระดับชั้น', 'คู่ที่/ทีมที่/มือวาง', 'สถานะ',
               'ผู้ยืนยัน/แก้ไข', 'เบอร์โทร', 'เวลาที่บันทึก', 'ชื่อเดิม', 'วันแข่ง', 'เวลารายงานตัว', 'ผู้รับรายงานตัว'])
    for e in entries:
        ws.append([e['event'], e['school'], e['prefix'], e['first'], e['last'], e['level'],
                   int(e['team']) if str(e['team']).isdigit() else e['team'],
                   e['status'] + (' (อัตโนมัติ เลยเวลาปิด)' if e.get('auto') else ''), e['by'], e['phone'], e['at'],
                   f"{e['oPrefix']}{e['oFirst']} {e['oLast']}", em.get(e['event'], {}).get('startDate', ''),
                   e['checkinAt'], e['checkinBy']])
    ws2 = wb.create_sheet('การยืนยัน')
    ws2.append(['โรงเรียน', 'รายการแข่งขัน', 'ผู้ยืนยัน', 'เบอร์โทร', 'เวลายืนยัน', 'จำนวนนักกีฬา'])
    for x in confs:
        ws2.append([x['school'], x['event'], x['by'], x['phone'], x['at'], x['count']])   # รวมรายการยืนยันอัตโนมัติ
    ws3 = wb.create_sheet('ประวัติการแก้ไข')
    ws3.append(['เวลา', 'โรงเรียน', 'รายการแข่งขัน', 'การกระทำ', 'ก่อน', 'หลัง', 'ผู้แก้ไข', 'เบอร์โทร', 'บัญชีเจ้าหน้าที่'])
    for r in get_db().execute('SELECT at, school, event, action, before, after, by_name, phone, staff_user FROM logs '
                              'WHERE comp_id=? ORDER BY id', (cid,)):
        ws3.append(list(r.values()))
    for sh in wb.worksheets:
        sh.freeze_panes = 'A2'
        for col in sh.columns:
            sh.column_dimensions[col[0].column_letter].width = min(max(len(str(x.value or '')) for x in col) + 2, 60)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{c['slug']}_{logic.now_bkk():%Y%m%d_%H%M}.xlsx",
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
